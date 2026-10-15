# -*- coding: utf-8 -*-
"""校验生成的 docx：包完整性、XML 合法性、关键元素次序、内容回读，并让 Word 实开一次。"""
import os
import sys
import zipfile
from xml.etree import ElementTree as ET

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# 每类父元素的合法子元素顺序（仅列出我们手工写入的节点相关的部分）
ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
            "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
            "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
            "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
            "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
            "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
            "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
            "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
            "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
            "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
            "eastAsianLayout", "specVanish", "oMath"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders",
             "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
             "hideMark"],
    "sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr",
               "type", "pgSz", "pgMar", "paperSrc", "pgBorders", "lnNumType",
               "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
               "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
               "printerSettings", "sectPrChange"],
    "settings": ["writeProtection", "view", "zoom", "removePersonalInformation",
                 "removeDateAndTime", "doNotDisplayPageBoundaries",
                 "displayBackgroundShape", "printPostScriptOverText",
                 "printFractionalCharacterWidth", "printFormsData",
                 "embedTrueTypeFonts", "embedSystemFonts", "saveSubsetFonts",
                 "saveFormsData", "mirrorMargins", "alignBordersAndEdges",
                 "bordersDoNotSurroundHeader", "bordersDoNotSurroundFooter",
                 "gutterAtTop", "hideSpellingErrors", "hideGrammaticalErrors",
                 "activeWritingStyle", "proofState", "formsDesign",
                 "attachedTemplate", "linkStyles", "stylePaneFormatFilter",
                 "stylePaneSortMethod", "documentType", "mailMerge",
                 "revisionView", "trackChanges", "doNotTrackMoves",
                 "doNotTrackFormatting", "documentProtection",
                 "autoFormatOverride", "styleLockTheme", "styleLockQFSet",
                 "defaultTabStop", "autoHyphenation", "consecutiveHyphenLimit",
                 "hyphenationZone", "doNotHyphenateCaps", "showEnvelope",
                 "summaryLength", "clickAndTypeStyle", "defaultTableStyle",
                 "evenAndOddHeaders", "bookFoldRevPrinting", "bookFoldPrinting",
                 "bookFoldPrintingSheets", "drawingGridHorizontalSpacing",
                 "drawingGridVerticalSpacing", "displayHorizontalDrawingGridEvery",
                 "displayVerticalDrawingGridEvery", "doNotUseMarginsForDrawingGridOrigin",
                 "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin",
                 "doNotShadeFormData", "noPunctuationKerning",
                 "characterSpacingControl", "printTwoOnOne",
                 "strictFirstAndLastChars", "noLineBreaksAfter",
                 "noLineBreaksBefore", "savePreviewPicture",
                 "doNotValidateAgainstSchema", "saveInvalidXml",
                 "ignoreMixedContent", "alwaysShowPlaceholderText",
                 "doNotDemarcateInvalidXml", "saveXmlDataOnly",
                 "useXSLTWhenSaving", "saveThroughXslt", "showXMLTags",
                 "alwaysMergeEmptyNamespace", "updateFields",
                 "hdrShapeDefaults", "footnotePr", "endnotePr", "compat",
                 "docVars", "rsids", "mathPr", "themeFontLang",
                 "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                 "doNotAutoCompressPictures", "forceUpgrade", "captions",
                 "readModeInkLockDown", "smartTagType", "schemaLibrary",
                 "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol",
                 "listSeparator"],
}


def check_order(root, problems):
    for tag, order in ORDER.items():
        for parent in root.iter(W + tag):
            pos = -1
            for child in parent:
                name = child.tag.split("}")[-1]
                if name not in order:
                    continue
                idx = order.index(name)
                if idx < pos:
                    problems.append(f"{tag}: {name} 次序错误")
                pos = idx


def main():
    path = os.path.abspath(sys.argv[1])
    print("文件:", path, os.path.getsize(path), "字节")
    problems = []

    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            problems.append(f"zip 损坏: {bad}")
        names = z.namelist()
        print("包内条目:", len(names))
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    data = z.read(name)
                    root = ET.fromstring(data)
                    if name.endswith("document.xml") or name.endswith("settings.xml"):
                        check_order(root, problems)
                except ET.ParseError as exc:
                    problems.append(f"{name} XML 解析失败: {exc}")

    print("XML 检查:", "通过" if not problems else f"{len(problems)} 处问题")
    for p in problems[:20]:
        print("  -", p)

    # 回读校验：确认内容完整、没有残留 Markdown 标记
    from docx import Document

    d = Document(path)
    paras = d.paragraphs
    text = "\n".join(p.text for p in paras)
    body_paras = [p for p in paras if p.text.strip()]
    styles = {}
    for p in body_paras:
        styles[p.style.name] = styles.get(p.style.name, 0) + 1
    headings = [(p.style.name, p.text) for p in body_paras
                if p.style.name.startswith("Heading")]
    print("回读: 段落", len(paras), "非空段", len(body_paras),
          "表格", len(d.tables), "字符", len(text))
    print("样式分布:", dict(sorted(styles.items(), key=lambda kv: -kv[1])[:8]))
    print("标题数:", len(headings), "| 前 3 条:", [h[1][:24] for h in headings[:3]])
    print("标题层级:", sorted({h[0] for h in headings}))
    leftovers = [t for t in ("**", "```", "| ---") if t in text]
    print("残留 Markdown 标记:", leftovers if leftovers else "无")

    # 用 Word 实开一次（只验证能否打开，不访问属性——本机属性访问被拒）
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(path, ReadOnly=True, AddToRecentFiles=False)
        print("Word 打开: 成功")
        try:
            doc.Close(SaveChanges=0)
        except Exception:
            pass
    except Exception as exc:
        print("Word 打开: 失败", repr(exc)[:200])
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
        os.system("taskkill /F /IM WINWORD.EXE >NUL 2>&1")


if __name__ == "__main__":
    main()
